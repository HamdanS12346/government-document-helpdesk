"""Spreadsheet-specific input processing boundary."""

from enum import StrEnum
from io import BytesIO
from typing import Protocol
from zipfile import BadZipFile

from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

from app.config.settings import Settings, get_settings
from app.contracts.normalized_input import (
    SpreadsheetCell,
    SpreadsheetContent,
    SpreadsheetMetadata,
    SpreadsheetSheet,
    SpreadsheetTable,
)
from app.input_processing.errors import InputProcessingError, InputProcessingErrorCode
from app.input_processing.schemas import (
    AttachmentProcessingError,
    InputModality,
    ValidatedAttachment,
)
from guardrails.input_processor import mark_document_text_untrusted, mask_pii_in_text

ParsedCellValue = str | int | float | bool | None


class SpreadsheetInspectionStatus(StrEnum):
    """Provider-independent spreadsheet inspection outcomes."""

    SUCCESS = "success"
    EXTRACTION_FAILURE = "extraction_failure"
    UNAVAILABLE = "unavailable"
    TIMEOUT = "timeout"
    MALFORMED_RESPONSE = "malformed_response"
    CORRUPT_WORKBOOK = "corrupt_workbook"
    UNSUPPORTED_PROTECTION = "unsupported_protection"
    WORKSHEET_LIMIT_EXCEEDED = "worksheet_limit_exceeded"
    PII_PROCESSING_FAILURE = "pii_processing_failure"


class ParsedWorksheet(BaseModel):
    """Provider-independent worksheet inspection summary."""

    model_config = ConfigDict(extra="forbid", strict=True)

    name: str
    position: int = Field(ge=1)
    visible: bool = True
    max_row: int = Field(ge=0)
    max_column: int = Field(ge=0)
    is_empty: bool = False
    cells: list["ParsedCell"] = Field(default_factory=list)
    merged_ranges: list[str] = Field(default_factory=list)
    tables: list["ParsedTable"] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)

    @field_validator("name")
    @classmethod
    def validate_name(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("worksheet name must not be empty")
        return value


class ParsedCell(BaseModel):
    """Provider-independent bounded cell extracted from a worksheet."""

    model_config = ConfigDict(extra="forbid", strict=True)

    coordinate: str
    row: int = Field(ge=1)
    column: int = Field(ge=1)
    value: ParsedCellValue = None
    value_type: str
    formula: str | None = None
    cached_value: ParsedCellValue = None
    truncated: bool = False
    untrusted_instruction_like: bool = False

    @field_validator("coordinate", "value_type")
    @classmethod
    def validate_non_empty_text(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("must not be empty")
        return value


class ParsedTable(BaseModel):
    """Provider-independent Excel Table metadata."""

    model_config = ConfigDict(extra="forbid", strict=True)

    name: str
    reference: str
    columns: list[str] = Field(default_factory=list)

    @field_validator("name", "reference")
    @classmethod
    def validate_non_empty_text(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("must not be empty")
        return value


class ParsedWorkbook(BaseModel):
    """Provider-independent workbook inspection result."""

    model_config = ConfigDict(extra="forbid", strict=True)

    filename: str
    worksheets: list[ParsedWorksheet] = Field(default_factory=list)

    @field_validator("filename")
    @classmethod
    def validate_filename(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("filename must not be empty")
        return value


class SpreadsheetInspectionResult(BaseModel):
    """Provider-independent spreadsheet parser result."""

    model_config = ConfigDict(extra="forbid", strict=True)

    status: SpreadsheetInspectionStatus
    workbook: ParsedWorkbook | None = None
    message: str | None = None

    @field_validator("message")
    @classmethod
    def validate_message(cls, value: str | None) -> str | None:
        if value is not None and not value.strip():
            raise ValueError("message must not be empty")
        return value

    @model_validator(mode="after")
    def validate_status_shape(self) -> "SpreadsheetInspectionResult":
        if self.status == SpreadsheetInspectionStatus.SUCCESS:
            if self.workbook is None:
                raise ValueError("successful spreadsheet inspection requires workbook")
            if self.message is not None:
                raise ValueError("successful spreadsheet inspection cannot include message")
            return self

        if self.workbook is not None:
            raise ValueError("failed spreadsheet inspection cannot include workbook")
        if self.message is None:
            raise ValueError("failed spreadsheet inspection requires a safe message")
        return self


class SpreadsheetProcessingResult(BaseModel):
    """Spreadsheet processor outcome with normalized content or safe failure."""

    model_config = ConfigDict(extra="forbid", strict=True)

    spreadsheet_content: SpreadsheetContent | None = None
    inspection_result: SpreadsheetInspectionResult | None = None
    error: AttachmentProcessingError | None = None

    @model_validator(mode="after")
    def validate_result_shape(self) -> "SpreadsheetProcessingResult":
        populated_fields = [
            self.spreadsheet_content is not None,
            self.inspection_result is not None,
            self.error is not None,
        ]
        if sum(populated_fields) != 1:
            raise ValueError(
                "spreadsheet processing result requires exactly one outcome field"
            )
        return self


class SpreadsheetParser(Protocol):
    """Narrow replaceable spreadsheet inspection capability."""

    def inspect(self, content: bytes, filename: str) -> SpreadsheetInspectionResult:
        """Inspect supported spreadsheet content from transient bytes."""


class PendingSpreadsheetParser:
    """Placeholder spreadsheet parser until the concrete provider is wired."""

    def inspect(self, content: bytes, filename: str) -> SpreadsheetInspectionResult:
        """Return a controlled unavailable outcome without exposing provider details."""

        return SpreadsheetInspectionResult(
            status=SpreadsheetInspectionStatus.UNAVAILABLE,
            message="Spreadsheet parser is not configured.",
        )


class OpenPyXLSpreadsheetParser:
    """OpenPyXL-backed workbook inspector behind the parser boundary."""

    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()

    def inspect(self, content: bytes, filename: str) -> SpreadsheetInspectionResult:
        """Inspect workbook structure from transient bytes without storing files."""

        workbook = None
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            workbook = load_workbook(
                BytesIO(content),
                read_only=False,
                data_only=False,
                keep_links=False,
            )
            cached_workbook = load_workbook(
                BytesIO(content),
                read_only=False,
                data_only=True,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.CORRUPT_WORKBOOK,
                message="Spreadsheet content could not be opened safely.",
            )
        except ImportError:
            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.UNAVAILABLE,
                message="Spreadsheet parser is not available.",
            )
        except Exception:
            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.EXTRACTION_FAILURE,
                message="Spreadsheet content could not be inspected.",
            )

        try:
            if _has_unsupported_workbook_protection(workbook):
                return SpreadsheetInspectionResult(
                    status=SpreadsheetInspectionStatus.UNSUPPORTED_PROTECTION,
                    message="Protected or encrypted workbooks are not supported.",
                )

            visible_worksheets = [
                (index, worksheet)
                for index, worksheet in enumerate(workbook.worksheets, start=1)
                if getattr(worksheet, "sheet_state", "visible") == "visible"
            ]
            visible_sheet_count = len(visible_worksheets)
            if visible_sheet_count > self.settings.spreadsheet_max_visible_sheets:
                return SpreadsheetInspectionResult(
                    status=SpreadsheetInspectionStatus.WORKSHEET_LIMIT_EXCEEDED,
                    message="Spreadsheet contains too many visible worksheets.",
                )
            worksheets = [
                _extract_worksheet_summary(
                    worksheet=worksheet,
                    cached_worksheet=cached_workbook[worksheet.title],
                    position=index,
                    settings=self.settings,
                )
                for index, worksheet in visible_worksheets
            ]

            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.SUCCESS,
                workbook=ParsedWorkbook(filename=filename, worksheets=worksheets),
            )
        except InputProcessingError as exc:
            if exc.code == InputProcessingErrorCode.PII_PROCESSING_FAILURE:
                return SpreadsheetInspectionResult(
                    status=SpreadsheetInspectionStatus.PII_PROCESSING_FAILURE,
                    message="Spreadsheet PII processing could not be completed safely.",
                )
            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.EXTRACTION_FAILURE,
                message="Spreadsheet content could not be inspected.",
            )
        except Exception:
            return SpreadsheetInspectionResult(
                status=SpreadsheetInspectionStatus.EXTRACTION_FAILURE,
                message="Spreadsheet content could not be inspected.",
            )
        finally:
            close = getattr(workbook, "close", None)
            if close is not None:
                close()
            cached_close = getattr(locals().get("cached_workbook", None), "close", None)
            if cached_close is not None:
                cached_close()


def _has_unsupported_workbook_protection(workbook: object) -> bool:
    security = getattr(workbook, "security", None)
    if security is None:
        return False

    protected_flags = (
        "lockStructure",
        "lockWindows",
        "lockRevision",
        "workbookPassword",
        "revisionsPassword",
    )
    return any(bool(getattr(security, flag, None)) for flag in protected_flags)


def _extract_worksheet_summary(
    *,
    worksheet: object,
    cached_worksheet: object,
    position: int,
    settings: Settings,
) -> ParsedWorksheet:
    raw_max_row = int(getattr(worksheet, "max_row", 0) or 0)
    raw_max_column = int(getattr(worksheet, "max_column", 0) or 0)

    bounded_max_row = min(raw_max_row, settings.spreadsheet_max_rows_per_sheet)
    bounded_max_column = min(
        raw_max_column,
        settings.spreadsheet_max_columns_per_sheet,
    )
    warnings = _limit_warnings(
        raw_max_row=raw_max_row,
        raw_max_column=raw_max_column,
        bounded_max_row=bounded_max_row,
        bounded_max_column=bounded_max_column,
    )

    cells = _extract_visible_cells(
        worksheet=worksheet,
        cached_worksheet=cached_worksheet,
        max_row=bounded_max_row,
        max_column=bounded_max_column,
        settings=settings,
    )
    warnings.extend(_cell_safety_warnings(cells))
    merged_ranges = _extract_merged_ranges(worksheet)
    tables = _extract_tables(worksheet)

    return ParsedWorksheet(
        name=worksheet.title,
        position=position,
        visible=True,
        max_row=max((cell.row for cell in cells), default=0),
        max_column=max((cell.column for cell in cells), default=0),
        is_empty=not cells,
        cells=cells,
        merged_ranges=merged_ranges,
        tables=tables,
        warnings=warnings,
    )


def _extract_visible_cells(
    *,
    worksheet: object,
    cached_worksheet: object,
    max_row: int,
    max_column: int,
    settings: Settings,
) -> list[ParsedCell]:
    cells: list[ParsedCell] = []
    hidden_rows = {
        index
        for index, dimension in worksheet.row_dimensions.items()
        if bool(getattr(dimension, "hidden", False))
    }
    hidden_columns = {
        index
        for index, dimension in worksheet.column_dimensions.items()
        if bool(getattr(dimension, "hidden", False))
    }

    for row in range(1, max_row + 1):
        if row in hidden_rows:
            continue
        for column in range(1, max_column + 1):
            cell = worksheet.cell(row=row, column=column)
            if _column_letter(column) in hidden_columns:
                continue
            if cell.value is None:
                continue
            formula = (
                cell.value
                if isinstance(cell.value, str) and cell.value.startswith("=")
                else None
            )
            cached_value = None
            value = cell.value
            if formula is not None:
                cached_cell = cached_worksheet.cell(row=row, column=column)
                cached_value = _sanitize_cell_value(
                    cached_cell.value,
                    max_characters=settings.spreadsheet_max_text_cell_characters,
                ).value
                value = None
            sanitized_value = _sanitize_cell_value(
                value,
                max_characters=settings.spreadsheet_max_text_cell_characters,
            )
            sanitized_formula = _sanitize_cell_value(
                formula,
                max_characters=settings.spreadsheet_max_text_cell_characters,
            )
            cells.append(
                ParsedCell(
                    coordinate=cell.coordinate,
                    row=row,
                    column=column,
                    value=sanitized_value.value,
                    value_type=_cell_value_type(cell.value),
                    formula=(
                        sanitized_formula.value
                        if isinstance(sanitized_formula.value, str)
                        else None
                    ),
                    cached_value=cached_value,
                    truncated=sanitized_value.truncated or sanitized_formula.truncated,
                    untrusted_instruction_like=(
                        sanitized_value.untrusted_instruction_like
                        or sanitized_formula.untrusted_instruction_like
                    ),
                )
            )
    return cells


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _extract_merged_ranges(worksheet: object) -> list[str]:
    return [
        str(merged_range)
        for merged_range in sorted(
            worksheet.merged_cells.ranges,
            key=lambda cell_range: (cell_range.min_row, cell_range.min_col),
        )
    ]


def _extract_tables(worksheet: object) -> list[ParsedTable]:
    tables = []
    for table in sorted(worksheet.tables.values(), key=lambda item: item.ref):
        columns = [
            str(column.name)
            for column in getattr(table, "tableColumns", [])
            if getattr(column, "name", None) is not None
        ]
        tables.append(
            ParsedTable(
                name=table.name,
                reference=table.ref,
                columns=columns,
            )
        )
    return tables


def _to_parsed_cell_value(value: object) -> ParsedCellValue:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()
    return str(value)


class SanitizedCellValue(BaseModel):
    """Bounded and masked cell value with safe metadata."""

    model_config = ConfigDict(extra="forbid", strict=True)

    value: ParsedCellValue
    truncated: bool = False
    untrusted_instruction_like: bool = False


def _sanitize_cell_value(
    value: object,
    *,
    max_characters: int,
) -> SanitizedCellValue:
    parsed_value = _to_parsed_cell_value(value)
    if not isinstance(parsed_value, str):
        return SanitizedCellValue(value=parsed_value)

    truncated = len(parsed_value) > max_characters
    bounded_value = parsed_value[:max_characters] if truncated else parsed_value
    masked_value = mask_pii_in_text(bounded_value).text
    untrusted_text = mark_document_text_untrusted(masked_value)
    return SanitizedCellValue(
        value=untrusted_text.text,
        truncated=truncated,
        untrusted_instruction_like=untrusted_text.suspicious,
    )


def _cell_value_type(value: object) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int | float):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, str) and value.startswith("#"):
        return "error"
    if hasattr(value, "isoformat"):
        return "date"
    return "string"


def _limit_warnings(
    *,
    raw_max_row: int,
    raw_max_column: int,
    bounded_max_row: int,
    bounded_max_column: int,
) -> list[str]:
    warnings = []
    if raw_max_row > bounded_max_row:
        warnings.append(
            f"Rows beyond configured limit were not extracted: {bounded_max_row} of {raw_max_row}."
        )
    if raw_max_column > bounded_max_column:
        warnings.append(
            f"Columns beyond configured limit were not extracted: {bounded_max_column} of {raw_max_column}."
        )
    return warnings


def _cell_safety_warnings(cells: list[ParsedCell]) -> list[str]:
    warnings = []
    truncated_coordinates = [
        cell.coordinate for cell in cells if cell.truncated
    ]
    if truncated_coordinates:
        warnings.append(
            "Text cell limit applied to "
            f"{len(truncated_coordinates)} cell(s): {', '.join(truncated_coordinates)}."
        )

    instruction_coordinates = [
        cell.coordinate for cell in cells if cell.untrusted_instruction_like
    ]
    if instruction_coordinates:
        warnings.append(
            "Instruction-like spreadsheet text was treated as untrusted data in "
            f"{len(instruction_coordinates)} cell(s): {', '.join(instruction_coordinates)}."
        )
    return warnings


def inspect_spreadsheet_workbook(
    *,
    content: bytes,
    filename: str,
    parser: SpreadsheetParser,
) -> SpreadsheetInspectionResult:
    """Inspect a workbook through a controlled parser boundary."""

    try:
        result = parser.inspect(content, filename)
    except Exception:
        return SpreadsheetInspectionResult(
            status=SpreadsheetInspectionStatus.EXTRACTION_FAILURE,
            message="Spreadsheet content could not be inspected.",
        )

    if not isinstance(result, SpreadsheetInspectionResult):
        return SpreadsheetInspectionResult(
            status=SpreadsheetInspectionStatus.MALFORMED_RESPONSE,
            message="Spreadsheet parser returned an invalid result.",
        )

    return result


def process_spreadsheet_attachment(
    validated_attachment: ValidatedAttachment,
    parser: SpreadsheetParser,
) -> SpreadsheetProcessingResult:
    """Process an already-validated spreadsheet through the parser boundary."""

    attachment = validated_attachment.attachment
    if validated_attachment.modality != InputModality.XLSX:
        return SpreadsheetProcessingResult(
            error=AttachmentProcessingError(
                filename=attachment.filename,
                code=InputProcessingErrorCode.UNSUPPORTED_FORMAT,
                message="This attachment is not a supported spreadsheet.",
            )
        )

    inspection = inspect_spreadsheet_workbook(
        content=attachment.content,
        filename=attachment.filename,
        parser=parser,
    )
    if inspection.status == SpreadsheetInspectionStatus.SUCCESS:
        if inspection.workbook is None:
            return SpreadsheetProcessingResult(
                error=AttachmentProcessingError(
                    filename=attachment.filename,
                    code=InputProcessingErrorCode.UNREADABLE_CONTENT,
                    message="Spreadsheet parser returned an invalid result.",
                )
            )
        return SpreadsheetProcessingResult(
            spreadsheet_content=build_spreadsheet_content(
                inspection.workbook,
                settings=get_settings(),
            )
        )

    return SpreadsheetProcessingResult(
        error=AttachmentProcessingError(
            filename=attachment.filename,
            code=_error_code_for_inspection_status(inspection.status),
            message=inspection.message or "Spreadsheet content could not be processed.",
        )
    )


def _error_code_for_inspection_status(
    status: SpreadsheetInspectionStatus,
) -> InputProcessingErrorCode:
    if status == SpreadsheetInspectionStatus.UNSUPPORTED_PROTECTION:
        return InputProcessingErrorCode.UNSUPPORTED_WORKBOOK_PROTECTION
    if status == SpreadsheetInspectionStatus.WORKSHEET_LIMIT_EXCEEDED:
        return InputProcessingErrorCode.SPREADSHEET_WORKSHEET_LIMIT_EXCEEDED
    if status == SpreadsheetInspectionStatus.PII_PROCESSING_FAILURE:
        return InputProcessingErrorCode.PII_PROCESSING_FAILURE
    if status in {
        SpreadsheetInspectionStatus.CORRUPT_WORKBOOK,
        SpreadsheetInspectionStatus.MALFORMED_RESPONSE,
    }:
        return InputProcessingErrorCode.UNREADABLE_CONTENT
    return InputProcessingErrorCode.EXTRACTION_FAILURE


def build_spreadsheet_content(
    workbook: ParsedWorkbook,
    *,
    settings: Settings,
) -> SpreadsheetContent:
    """Build normalized spreadsheet content from provider-independent parsed data."""

    sheets = [
        SpreadsheetSheet(
            name=sheet.name,
            position=sheet.position,
            max_row=sheet.max_row,
            max_column=sheet.max_column,
            is_empty=sheet.is_empty,
            cells=[
                SpreadsheetCell(
                    coordinate=cell.coordinate,
                    row=cell.row,
                    column=cell.column,
                    value=cell.value,
                    value_type=cell.value_type,
                    formula=cell.formula,
                    cached_value=cell.cached_value,
                    truncated=cell.truncated,
                )
                for cell in sheet.cells
            ],
            merged_ranges=sheet.merged_ranges,
            tables=[
                SpreadsheetTable(
                    name=table.name,
                    reference=table.reference,
                    columns=table.columns,
                )
                for table in sheet.tables
            ],
        )
        for sheet in workbook.worksheets
    ]
    warnings = [
        warning
        for sheet in workbook.worksheets
        for warning in sheet.warnings
    ]
    metadata = SpreadsheetMetadata(
        workbook_name=workbook.filename,
        processed_sheet_count=len(sheets),
        total_visible_sheet_count=len(sheets),
        hidden_sheet_count=0,
        max_sheets=settings.spreadsheet_max_visible_sheets,
        max_rows_per_sheet=settings.spreadsheet_max_rows_per_sheet,
        max_columns_per_sheet=settings.spreadsheet_max_columns_per_sheet,
        max_text_cell_characters=settings.spreadsheet_max_text_cell_characters,
        preview_row_count=settings.spreadsheet_preview_row_count,
    )
    content = SpreadsheetContent(
        workbook_name=workbook.filename,
        sheets=sheets,
        preview="",
        warnings=warnings,
        metadata=metadata,
    )
    return content.model_copy(
        update={"preview": build_spreadsheet_preview(content)}
    )


def build_spreadsheet_preview(content: SpreadsheetContent) -> str:
    """Build deterministic bounded preview text for a workbook."""

    lines = [f"Workbook: {content.workbook_name}", "Sheets:"]
    for sheet in content.sheets:
        status = "empty" if sheet.is_empty else f"{sheet.max_row} rows x {sheet.max_column} columns"
        lines.append(f"- {sheet.position}. {sheet.name}: {status}")

    for sheet in content.sheets:
        lines.append("")
        lines.append(f"Sheet: {sheet.name}")
        if sheet.is_empty:
            lines.append("[empty sheet]")
            continue
        sample_rows = _sample_sheet_rows(
            sheet,
            row_limit=content.metadata.preview_row_count,
        )
        lines.extend(sample_rows)
    return "\n".join(lines)


def build_spreadsheet_text_projection(content: SpreadsheetContent) -> str:
    """Project normalized spreadsheet content into deterministic plain text."""

    lines = [f"Workbook: {content.workbook_name}"]
    for sheet in content.sheets:
        lines.append("")
        lines.append(f"Sheet {sheet.position}: {sheet.name}")
        if sheet.is_empty:
            lines.append("[empty sheet]")
        else:
            lines.append(f"Dimensions: {sheet.max_row} rows x {sheet.max_column} columns")
            lines.extend(_sample_sheet_rows(sheet, row_limit=sheet.max_row))
        if sheet.merged_ranges:
            lines.append(f"Merged ranges: {', '.join(sheet.merged_ranges)}")
        for table in sheet.tables:
            columns = ", ".join(table.columns) if table.columns else "[no columns]"
            lines.append(
                f"Table: {table.name} ({table.reference}) columns: {columns}"
            )
    if content.warnings:
        lines.append("")
        lines.append("Warnings:")
        lines.extend(f"- {warning}" for warning in content.warnings)
    return "\n".join(lines)


def _sample_sheet_rows(sheet: SpreadsheetSheet, *, row_limit: int) -> list[str]:
    rows: dict[int, list[SpreadsheetCell]] = {}
    for cell in sheet.cells:
        rows.setdefault(cell.row, []).append(cell)

    lines = []
    for row_number in sorted(rows)[:row_limit]:
        cells = sorted(rows[row_number], key=lambda cell: cell.column)
        rendered_cells = [
            f"{cell.coordinate}={_render_cell_value(cell)}"
            for cell in cells
        ]
        lines.append(f"Row {row_number}: " + " | ".join(rendered_cells))
    return lines


def _render_cell_value(cell: SpreadsheetCell) -> str:
    if cell.formula is not None:
        rendered = f"formula {cell.formula}"
        if cell.cached_value is not None:
            rendered += f" cached {cell.cached_value}"
        return rendered
    if cell.value is None:
        return "[blank]"
    return str(cell.value)


__all__ = [
    "OpenPyXLSpreadsheetParser",
    "ParsedCell",
    "ParsedCellValue",
    "ParsedTable",
    "ParsedWorkbook",
    "ParsedWorksheet",
    "PendingSpreadsheetParser",
    "SpreadsheetInspectionResult",
    "SpreadsheetInspectionStatus",
    "SpreadsheetParser",
    "SpreadsheetProcessingResult",
    "build_spreadsheet_content",
    "build_spreadsheet_preview",
    "build_spreadsheet_text_projection",
    "inspect_spreadsheet_workbook",
    "process_spreadsheet_attachment",
]
